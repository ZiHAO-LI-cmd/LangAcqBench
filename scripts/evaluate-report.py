#!/usr/bin/env python3
"""Combine translation and lm_eval JSON results in a readable Excel workbook."""

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADERS = ('Task', 'Dataset / Direction', 'Samples', 'BLEU', 'chrF++', 'Acc_norm', 'Exact match')
DATASET_NAMES = {'bouquet': 'Bouquet', 'flores200': 'FLORES-200'}
PIQA_VARIANTS = {
    'nonparallel_cloze': 'Nonparallel cloze',
    'nonparallel_generation': 'Nonparallel generation',
    'parallel_cloze': 'Parallel cloze',
    'parallel_generation': 'Parallel generation',
}


def read_json(path):
    with path.open(encoding='utf-8') as source:
        return json.load(source)


def language_name(code, languages):
    if code in languages:
        return languages[code]
    for language_code, name in languages.items():
        if code.casefold() == language_code.casefold() or code.split('_')[0].casefold() == language_code.split('_')[0].casefold():
            return name
    return code


def metric(result, prefix):
    values = [value for key, value in result.items() if key.startswith(prefix + ',')]
    if len(values) > 1:
        raise ValueError(f'Ambiguous {prefix} metric for {result.get("name", "task")}')
    return values[0] if values else None


def task_label(task, languages, config):
    language = language_name(config.get('dataset_name') or task.split('_')[-1], languages)
    if task.startswith('belebele_'):
        return 'Belebele', language
    if task.startswith('multiblimp_'):
        return 'MultiBLiMP', language
    if task.startswith('global_piqa_'):
        suffix = task[len('global_piqa_'):]
        for variant, label in PIQA_VARIANTS.items():
            if suffix.startswith(variant + '_'):
                return 'Global PIQA', f'{label}, {language}'
    return task, language


def build_rows(translation, lm_results=None):
    languages = translation['languages']
    rows = []
    for group in translation['groups']:
        direction = f'{language_name(group["source"], languages)} → {language_name(group["target"], languages)}'
        dataset = DATASET_NAMES.get(group['dataset'].casefold(), group['dataset'])
        rows.append(('Translation', f'{dataset}: {direction}', group['count'], group['bleu'], group['chrf++'], None, None))

    if lm_results is not None:
        results = lm_results['results']
        configs = lm_results.get('configs', {})
        samples = lm_results.get('n-samples', {})
        for task, result in results.items():
            label, dataset = task_label(task, languages, configs.get(task, {}))
            count = samples.get(task, {}).get('effective', result.get('sample_len'))
            rows.append((label, dataset, count, None, None,
                         metric(result, 'acc_norm'), metric(result, 'exact_match')))
    averages = []
    for column in (3, 4, 5, 6):
        values = [row[column] for row in rows if row[column] is not None]
        averages.append(sum(values) / len(values) if values else None)
    return [('Avg', None, None, *averages)] + rows


def write_workbook(rows, output):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Evaluation'
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)

    dark = PatternFill('solid', fgColor='17365D')
    light = PatternFill('solid', fgColor='DDEBF7')
    for cell in sheet[1]:
        cell.fill = dark
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center')
    for cell in sheet[2]:
        cell.fill = light
        cell.font = Font(bold=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        row[2].number_format = '#,##0'
        for cell in row[3:5]:
            cell.number_format = '0.00'
        for cell in row[5:7]:
            cell.number_format = '0.00%'
    for index, width in enumerate((18, 43, 12, 12, 12, 15, 16), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = 'A3'
    sheet.auto_filter.ref = f'A1:G{sheet.max_row}'
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--translation', type=Path, required=True, help='Translation evaluator JSON')
    parser.add_argument('--lm-eval-dir', type=Path, help='Directory containing one lm-eval_*.json result')
    parser.add_argument('--output', type=Path, required=True, help='Excel workbook path')
    args = parser.parse_args()
    try:
        translation = read_json(args.translation)
        lm_results = None
        if args.lm_eval_dir is not None:
            files = sorted(args.lm_eval_dir.glob('lm-eval_*.json'))
            if len(files) != 1:
                raise ValueError(f'Expected one lm_eval result in {args.lm_eval_dir}; found {len(files)}')
            lm_results = read_json(files[0])
        write_workbook(build_rows(translation, lm_results), args.output)
        print(f'Excel report: {args.output}')
    except (OSError, KeyError, TypeError, ValueError) as exc:
        parser.exit(2, f'evaluate-report: {exc}\n')


if __name__ == '__main__':
    main()
